"""Build the IEEE paper (DOCX on the prescribed ProComm template, Markdown, and PDF via Word).

    python reports/build_paper.py            # DOCX + Markdown (+ PDF and page previews on Windows with Word)
    python reports/build_paper.py --no-pdf   # skip the Word export

Content and numbers come from reports/paper_content.py, which reads the saved
experiment outputs. Citations are numbered in order of first appearance.
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "reports"))

from docx import Document  # noqa: E402
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from lxml import etree  # noqa: E402

from ckd_xai.report_data import load_results  # noqa: E402
from paper_content import EQUATIONS, REFERENCES, build_blocks  # noqa: E402

TEMPLATE = ROOT / "reports" / "templates" / "Full-Paper-template.docx"
OUT_DOCX = ROOT / "reports" / "IEEE_paper_CKD_XAI.docx"
OUT_PDF = ROOT / "reports" / "IEEE_paper_CKD_XAI.pdf"
OUT_MD = ROOT / "reports" / "IEEE_paper_CKD_XAI.md"
PREVIEW_DIR = ROOT / "reports" / "build"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
COL_WIDTH_IN = 3.25
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
TABLE_PT = 9


# --------------------------------------------------------------------------- numbering
def number_items(content):
    cites, figs, tabs, eqs = {}, {}, {}, {}

    def scan(text):
        for m in re.finditer(r"\[@([^\]]+)\]", text):
            for k in m.group(1).split(";"):
                k = k.strip()
                if k not in REFERENCES:
                    raise KeyError(f"Unknown citation key: {k}")
                cites.setdefault(k, len(cites) + 1)

    for b in content["blocks"]:
        kind = b[0]
        if kind == "p":
            scan(b[1])
        elif kind in ("numbered", "bullets"):
            for t in b[1]:
                scan(t)
        elif kind == "fig":
            figs[b[2]] = len(figs) + 1
            scan(b[3])
        elif kind == "table":
            tabs[b[1]] = len(tabs) + 1
            scan(b[2])
            for row in b[4]:
                for cell in row:
                    scan(cell)
        elif kind == "eq":
            eqs[b[1]] = len(eqs) + 1
    unused = set(REFERENCES) - set(cites)
    if unused:
        raise ValueError(f"References never cited: {sorted(unused)}")
    return cites, figs, tabs, eqs


def format_cites(keys, cites):
    nums = sorted(cites[k.strip()] for k in keys)
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            parts.append(f"[{nums[i]}]–[{nums[j]}]")
        else:
            parts.extend(f"[{n}]" for n in nums[i:j + 1])
        i = j + 1
    return ", ".join(parts)


def resolve(text, num):
    cites, figs, tabs, eqs = num
    text = re.sub(r"\[@([^\]]+)\]", lambda m: format_cites(m.group(1).split(";"), cites), text)
    text = re.sub(r"\{fig:(\w+)\}", lambda m: f"Figure {figs[m.group(1)]}", text)
    text = re.sub(r"\{tab:(\w+)\}", lambda m: f"Table {tabs[m.group(1)]}", text)
    text = re.sub(r"\{eq:(\w+)\}", lambda m: f"({eqs[m.group(1)]})", text)
    return text


TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")


def inline_runs(text):
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], {"b"}))
        elif part.startswith("*"):
            out.append((part[1:-1], {"i"}))
        elif part.startswith("_{"):
            out.append((part[2:-1], {"sub"}))
        elif part.startswith("^{"):
            out.append((part[2:-1], {"sup"}))
        else:
            out.append((part, set()))
    return out


# --------------------------------------------------------------------------- DOCX helpers
def add_runs(par, text, size=None, italic=None):
    lines = text.split("\n")
    if len(lines) > 1:                      # explicit line breaks (used inside table cells)
        for i, line in enumerate(lines):
            if i:
                par.add_run().add_break()
            add_runs(par, line, size, italic)
        return par
    for chunk, fmt in inline_runs(text):
        r = par.add_run(chunk)
        if "b" in fmt:
            r.bold = True
        if "i" in fmt:
            r.italic = True
        if "sub" in fmt:
            r.font.subscript = True
        if "sup" in fmt:
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)
        if italic is not None and "i" not in fmt:
            r.italic = italic
    return par


def omml_for(key):
    body = EQUATIONS[key][1]
    if not MML2OMML.exists():
        return None
    mml = etree.fromstring(f'<math xmlns="http://www.w3.org/1998/Math/MathML">{body}</math>')
    xslt = etree.XSLT(etree.parse(str(MML2OMML)))
    return xslt(mml).getroot()


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, (val, sz) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)


def table_borders(table, none=False):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    outer = "nil" if none else "single"
    for edge, val in (("top", outer), ("bottom", outer), ("left", "nil"), ("right", "nil"),
                      ("insideH", "nil"), ("insideV", "nil")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        if val == "single":
            el.set(qn("w:sz"), "8")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        b.append(el)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), "43")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


class PaperDocx:
    def __init__(self, num, R):
        self.num = num
        self.doc = Document(str(TEMPLATE))
        self.body = self.doc.element.body
        self.final_sectPr = self.body.find(qn("w:sectPr"))
        self._strip_template()

    # ---- template preparation
    def _strip_template(self):
        body = self.body
        for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
            for el in list(body.iter(qn(tag))):
                el.getparent().remove(el)
        for ref in list(body.iter(qn("w:commentReference"))):
            run = ref.getparent()
            run.getparent().remove(run)
        children = list(body)
        first_break = next(i for i, el in enumerate(children)
                           if el.tag == qn("w:p") and el.find(".//" + qn("w:sectPr")) is not None)
        for el in children[first_break + 1:]:
            if el is not self.final_sectPr:
                body.remove(el)
        for rid, rel in list(self.doc.part.rels.items()):
            if any(s in rel.reltype for s in ("/comments", "commentsExtended", "commentsIds", "/people")):
                self.doc.part.drop_rel(rid)

    def title_block(self, title, author):
        tbl = self.doc.tables[0]
        tp = tbl.rows[0].cells[0].paragraphs[0]
        for r in list(tp._p.iter(qn("w:r"))):
            r.getparent().remove(r)
        tp.add_run(title)
        row = tbl.rows[1]
        cell = row.cells[0].merge(row.cells[2])
        paras = cell.paragraphs
        for p in paras[1:]:
            p._p.getparent().remove(p._p)
        p = paras[0]
        for child in list(p._p):
            if child.tag != qn("w:pPr"):
                p._p.remove(child)
        p.add_run(author["name"])
        for line in author["affiliation"]:
            p.add_run().add_break()
            p.add_run(line).italic = True

    # ---- blocks
    def para(self, text, style="Body Text Indent", indent=True):
        p = self.doc.add_paragraph(style=style)
        add_runs(p, resolve(text, self.num))
        if not indent:
            p.paragraph_format.first_line_indent = Inches(0)
        return p

    def abstract(self, abstract, index_terms):
        p = self.doc.add_paragraph(style="Abstract")
        p.add_run("Abstract - " + abstract)
        self.doc.add_paragraph(style="Abstract")
        p = self.doc.add_paragraph(style="Abstract")
        p.add_run("Index Terms - " + index_terms)

    def heading(self, text, level):
        self.doc.add_paragraph(text, style=f"Heading {level}")

    def listing(self, items, style):
        for it in items:
            self.doc.add_paragraph(style=style)
            add_runs(self.doc.paragraphs[-1], resolve(it, self.num))

    def equation(self, key):
        """Display-style equation, centred, with its number flush right (borderless 1x2 table)."""
        n = self.num[3][key]
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_borders(t, none=True)
        left, right = t.rows[0].cells
        left.width, right.width = Inches(COL_WIDTH_IN - 0.4), Inches(0.4)
        for c in (left, right):
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            pf = c.paragraphs[0].paragraph_format
            pf.space_before, pf.space_after = Pt(3), Pt(3)
        om = omml_for(key)
        if om is not None:
            para = OxmlElement("m:oMathPara")
            ppr = OxmlElement("m:oMathParaPr")
            jc = OxmlElement("m:jc")
            jc.set(qn("m:val"), "center")
            ppr.append(jc)
            para.append(ppr)
            para.append(om)
            left.paragraphs[0]._p.append(para)
        else:
            left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            left.paragraphs[0].add_run(re.sub(r"<[^>]+>", "", EQUATIONS[key][1]))
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraphs[0].add_run(f"({n})")

    def figure(self, path, key, caption, width=COL_WIDTH_IN):
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(ROOT / path), width=Inches(width))
        cap = self.doc.add_paragraph(style="Caption")
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.keep_together = True
        cap.add_run(f"Figure {self.num[1][key]}. ")
        add_runs(cap, resolve(caption, self.num))

    def section_break(self, cols):
        sp = copy.deepcopy(self.final_sectPr)
        c = sp.find(qn("w:cols"))
        for a in list(c.attrib):
            del c.attrib[a]
        if cols == 2:
            c.set(qn("w:num"), "2")
            c.set(qn("w:space"), "461")
        else:
            c.set(qn("w:space"), "720")
        p = self.doc.add_paragraph(style="Normal")
        pPr = p._p.get_or_add_pPr()
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "2")
        rpr.append(sz)
        pPr.append(rpr)
        pPr.append(sp)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)

    def table(self, key, title, header, rows, widths, full, font_pt=TABLE_PT):
        if full:
            self.section_break(cols=2)          # end the two-column section above the table
        cap = self.doc.add_paragraph(style="Caption")
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.keep_together = True
        if full:
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run(f"Table {self.num[2][key]}. ")
        add_runs(cap, resolve(title, self.num))
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_borders(t)
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]
            c.paragraphs[0].style = self.doc.styles["Table Head"]
            add_runs(c.paragraphs[0], h, size=font_pt)
            set_cell_border(c, bottom=("single", 4))
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                c = t.rows[i].cells[j]
                c.paragraphs[0].style = self.doc.styles["Table text"]
                add_runs(c.paragraphs[0], resolve(val, self.num), size=font_pt)
        text_cols = {"features": {0, 1}, "shap": {1, 4}}.get(key, {0})
        for row in t.rows:
            for j, w in enumerate(widths):
                cell = row.cells[j]
                cell.width = Inches(w)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.LEFT if j in text_cols
                                                else WD_ALIGN_PARAGRAPH.CENTER)
        # Keep the whole table in one column: rows may not split, and every row stays with the next.
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            if row._tr is not t.rows[-1]._tr:
                for c in row.cells:
                    for par in c.paragraphs:
                        par.paragraph_format.keep_with_next = True
        spacer = self.doc.add_paragraph(style="Normal")
        spacer.paragraph_format.space_after = Pt(4)
        if full:
            self.section_break(cols=1)          # end the one-column section holding the table

    def references(self, cites):
        self.heading("References", 1)
        for key, n in sorted(cites.items(), key=lambda kv: kv[1]):
            p = self.doc.add_paragraph(style="Heading 3")
            pf = p.paragraph_format
            pf.left_indent = Inches(0.3)
            pf.first_line_indent = Inches(-0.3)
            pf.space_after = Pt(2)
            pf.tab_stops.add_tab_stop(Inches(0.3))
            p.add_run(f"[{n}]\t")
            add_runs(p, REFERENCES[key])

    def finish(self, title, balance=True):
        if balance:
            self.section_break(cols=2)          # balances the columns on the last page
        cp = self.doc.core_properties
        cp.title = title
        cp.author = "Vedank Singh"
        cp.subject = "IEEE-format mini-project report"
        cp.keywords = "chronic kidney disease; explainable AI; SHAP; machine learning"
        cp.comments = ""
        cp.last_modified_by = "Vedank Singh"
        self.doc.save(str(OUT_DOCX))


# --------------------------------------------------------------------------- build
def build_docx(content, num, balance=True):
    d = PaperDocx(num, None)
    d.title_block(content["title"], content["author"])
    d.abstract(content["abstract"], content["index_terms"])
    sec, sub = 0, 0
    for b in content["blocks"]:
        kind = b[0]
        if kind == "h1":
            sec += 1
            sub = 0
            d.heading(f"{ROMAN[sec - 1]}. {b[1]}", 1)
        elif kind == "h1u":
            d.heading(b[1], 1)
        elif kind == "h2":
            sub += 1
            d.heading(f"{chr(64 + sub)}. {b[1]}", 2)
        elif kind == "p":
            d.para(b[1], indent=not (len(b) > 2 and b[2] == "noindent"))
        elif kind == "numbered":
            d.listing(b[1], "NumberedList")
        elif kind == "bullets":
            d.listing(b[1], "BulletList")
        elif kind == "eq":
            d.equation(b[1])
        elif kind == "fig":
            d.figure(b[1], b[2], b[3], *(b[4:5] or []))
        elif kind == "table":
            d.table(*b[1:])
        elif kind == "refs":
            d.references(num[0])
    d.finish(content["title"], balance)


def md_inline(text, num):
    text = resolve(text, num)
    text = re.sub(r"_\{(.+?)\}", r"<sub>\1</sub>", text)
    text = re.sub(r"\^\{(.+?)\}", r"<sup>\1</sup>", text)
    return text


def build_markdown(content, num):
    cites, figs, tabs, eqs = num
    L = [f"# {content['title']}", "", f"**{content['author']['name']}**  "]
    L += [f"*{a}*  " for a in content["author"]["affiliation"]]
    L += ["", f"**Abstract** - {content['abstract']}", "", f"**Index Terms** - {content['index_terms']}", ""]
    sec = sub = 0
    for b in content["blocks"]:
        kind = b[0]
        if kind == "h1":
            sec += 1
            sub = 0
            L += [f"## {ROMAN[sec - 1]}. {b[1]}", ""]
        elif kind == "h1u":
            L += [f"## {b[1]}", ""]
        elif kind == "h2":
            sub += 1
            L += [f"### {chr(64 + sub)}. {b[1]}", ""]
        elif kind == "p":
            L += [md_inline(b[1], num), ""]
        elif kind == "numbered":
            L += [f"{i}) {md_inline(t, num)}" for i, t in enumerate(b[1], 1)] + [""]
        elif kind == "bullets":
            L += [f"- {md_inline(t, num)}" for t in b[1]] + [""]
        elif kind == "eq":
            L += [f"$$ {EQUATIONS[b[1]][0]} \\tag{{{eqs[b[1]]}}} $$", ""]
        elif kind == "fig":
            rel = "../" + b[1]
            L += [f"![Figure {figs[b[2]]}]({rel})", "", f"*Figure {figs[b[2]]}. {md_inline(b[3], num)}*", ""]
        elif kind == "table":
            key, title, header, rows = b[1], b[2], b[3], b[4]
            L += [f"**Table {tabs[key]}.** {md_inline(title, num)}", ""]
            L += ["| " + " | ".join(header) + " |", "|" + "|".join("---" for _ in header) + "|"]
            L += ["| " + " | ".join(md_inline(c, num) for c in r) + " |" for r in rows] + [""]
        elif kind == "refs":
            L += ["## References", ""]
            for k, n in sorted(cites.items(), key=lambda kv: kv[1]):
                L += [f"[{n}] {REFERENCES[k]}", ""]
    OUT_MD.write_text("\n".join(L), encoding="utf-8")


def export_pdf():
    ps = (
        "$ErrorActionPreference='Stop'; $w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{OUT_DOCX}', $false, $true); "
        f"$d.ExportAsFixedFormat('{OUT_PDF}', 17); "
        "$pages = $d.ComputeStatistics(2); $d.Close($false); $w.Quit(); Write-Output $pages"
    )
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=300)
    if res.returncode != 0:
        raise RuntimeError(res.stderr)
    return res.stdout.strip()


def last_page_blank():
    import pymupdf

    doc = pymupdf.open(str(OUT_PDF))
    last = doc[-1]
    return not last.get_text().strip() and not last.get_images()


def previews():
    import pymupdf as fitz

    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    for old in PREVIEW_DIR.glob("paper_page_*.png"):
        old.unlink()
    doc = fitz.open(str(OUT_PDF))
    for i, page in enumerate(doc, 1):
        page.get_pixmap(dpi=110).save(str(PREVIEW_DIR / f"paper_page_{i:02d}.png"))
    return len(doc)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()
    R = load_results()
    content = build_blocks(R)
    num = number_items(content)
    build_docx(content, num)
    build_markdown(content, num)
    print(f"wrote {OUT_DOCX.name} and {OUT_MD.name}; {len(num[0])} references, {len(num[1])} figures, "
          f"{len(num[2])} tables, {len(num[3])} equations")
    if not args.no_pdf:
        pages = export_pdf()
        if last_page_blank():
            # The empty section that balances the final columns spilled onto a new page: drop it.
            build_docx(content, num, balance=False)
            pages = export_pdf()
            print("note: last page was full, so the column-balancing break was omitted")
        n = previews()
        print(f"wrote {OUT_PDF.name}: Word reports {pages} pages; {n} preview images in reports/build/")


if __name__ == "__main__":
    main()
