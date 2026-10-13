"""Fill the college front page (kept separate from the IEEE paper) and make a combined submission PDF.

    python reports/build_front_page.py

Only details stated in the project synopsis are filled in. The roll number is not
in the synopsis, so that cell is left blank for the student to complete.
"""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "reports" / "templates" / "ML_Front_Page.docx"
OUT_DOCX = ROOT / "reports" / "Front_Page_CKD_XAI.docx"
OUT_PDF = ROOT / "reports" / "Front_Page_CKD_XAI.pdf"
PAPER_PDF = ROOT / "reports" / "IEEE_paper_CKD_XAI.pdf"
COMBINED = ROOT / "reports" / "Submission_FrontPage_and_IEEE_Paper.pdf"

TITLE = "Explainable Machine Learning for Chronic Kidney Disease Prediction Using SHAP-Based Feature Attribution"
STUDENT = {"Name": "Vedank Singh", "Registration No": "240962043", "Roll No": "",
           "Semester & Branch": "V Semester, B.Tech. CSE (AI & ML)", "Section": "B"}
GUIDES = "Under the guidance of Dr. Narendra V G, Professor, and Ms Vibha, Assistant Professor - Senior Scale"
DATE_LINE = "October 2026 (Academic Year 2026–27)"


def set_text(paragraph, text):
    """Replace the text but keep the first run's formatting."""
    runs = paragraph.runs
    first = next((r for r in runs if r.text.strip()), runs[0])
    for r in runs:
        if r is not first:
            r._r.getparent().remove(r._r)
    first.text = text


def main():
    d = Document(str(TEMPLATE))
    # The real title wraps to three lines, so drop some of the template's blank spacer lines
    # (keeping one in each gap) to keep the cover on a single page.
    paras = d.paragraphs
    texts = [p.text.strip() for p in paras]
    i_course = texts.index("Machine Learning (CSE 3125)")
    i_title = texts.index("Project Title")
    i_guide = next(i for i, t in enumerate(texts) if t.startswith("Under the guidance"))
    i_school = texts.index("School of Computer Engineering")
    blanks = [i for i in range(i_course + 1, i_title) if not texts[i]][1:]
    blanks += [i for i in range(i_guide + 1, i_school) if not texts[i]][1:]
    for i in sorted(blanks, reverse=True):
        paras[i]._p.getparent().remove(paras[i]._p)
    for p in d.paragraphs:
        t = p.text.strip()
        if t == "Project Title":
            set_text(p, TITLE)
        elif t.startswith("Under the guidance of"):
            set_text(p, GUIDES)
        elif t.startswith("October"):
            set_text(p, DATE_LINE)
    table = d.tables[0]
    header = [c.text.strip() for c in table.rows[0].cells]
    row = table.rows[1]
    for key, value in STUDENT.items():
        cell = row.cells[header.index(key)]
        para = cell.paragraphs[0]
        if para.runs:
            set_text(para, value)
        else:
            para.add_run(value)
    for extra in list(table.rows)[2:]:           # one-person project: drop the unused member rows
        extra._tr.getparent().remove(extra._tr)
    cp = d.core_properties
    cp.author, cp.title = "Vedank Singh", f"Front page - {TITLE}"
    d.save(str(OUT_DOCX))

    ps = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT_DOCX}', $false, $true); $d.ExportAsFixedFormat('{OUT_PDF}', 17); "
          "$d.Close($false); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=300)
    import pymupdf

    combined = pymupdf.open()
    for src in (OUT_PDF, PAPER_PDF):
        with pymupdf.open(str(src)) as part:
            combined.insert_pdf(part)
    combined.save(str(COMBINED))
    print(f"wrote {OUT_DOCX.name}, {OUT_PDF.name} and {COMBINED.name} ({combined.page_count} pages)")
    page = pymupdf.open(str(OUT_PDF))[0]
    page.get_pixmap(dpi=80).save(str(ROOT / "reports" / "build" / "front_page_preview.png"))


if __name__ == "__main__":
    main()
